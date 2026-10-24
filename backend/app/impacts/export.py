"""Impact-study exports: long-format CSV and a multi-sheet Excel workbook (JSON is the stored result itself).

Every row carries its model-status label, so an exported number is never separated from how it was obtained.
"""

from __future__ import annotations

import csv
import io

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from app.impacts import catalogue
from app.impacts.costing import EXCLUDED_GLOBAL

LIMITATIONS = [
    "Screening-level decision support; not validated for construction and not certified for regulatory compliance.",
    "Combinations are simulated as whole configurations; the non-additive remainder is a diagnostic, not a causal attribution.",
    "CAPEX uses only unit rates entered for the project (source and basis year shown per rate); missing rates are reported as data gaps.",
    "Space results are preliminary planning ratios and configured geometry, not architectural measurements.",
    "Embodied carbon covers added strategy materials only, and only where emission factors were supplied; the base building is excluded.",
    "Natural-ventilation and earth-air results come from lumped models (zonal airflow network, radial FD soil) that require CFD and "
    "site investigation before design.",
]


def scenarios(result: dict) -> list[dict]:
    out = [result["baseline"]] + result.get("individual", [])
    if result.get("combination"):
        out.append(result["combination"])
    return out


def long_rows(result: dict) -> list[dict]:
    rows = []
    for ev in scenarios(result):
        for r in ev.get("metrics", []):
            dr = r.get("delta_range") or [None, None]
            rows.append({"scenario": ev["id"], "scenario_label": ev["label"], "strategies": "+".join(ev.get("resolved", [])),
                         "auto_added": "+".join(ev.get("auto_added", [])), "feasible": ev.get("feasible"),
                         "infeasible_reasons": "; ".join(ev.get("infeasible_reasons", [])), "metric": r["key"], "metric_label": r["label"],
                         "group": r["group"], "unit": r["unit"], "baseline": r["baseline"], "value": r["value"], "delta": r["delta"],
                         "delta_low": dr[0], "delta_high": dr[1], "pct_change": r.get("pct"), "pct_note": r.get("pct_note"),
                         "status": r["status"], "note": r.get("note", "")})
    for p in result.get("pairs", []):
        for key, m in p["metrics"].items():
            rows.append({"scenario": f"{p['a']}+{p['b']}", "scenario_label": f"{p['a']} + {p['b']}", "strategies": f"{p['a']}+{p['b']}",
                         "auto_added": "+".join(p.get("auto_added", [])), "feasible": p["feasible"],
                         "infeasible_reasons": "; ".join(p["infeasible_reasons"]), "metric": key, "metric_label": key, "group": "",
                         "unit": "", "baseline": None, "value": m["value"], "delta": m["delta"], "delta_low": None, "delta_high": None,
                         "pct_change": m["pct"], "pct_note": None, "status": m["status"], "note": ""})
    return rows


def to_csv(result: dict) -> str:
    rows = long_rows(result)
    buf = io.StringIO()
    w = csv.DictWriter(buf, fieldnames=list(rows[0].keys()) if rows else ["scenario"])
    w.writeheader()
    w.writerows(rows)
    return buf.getvalue()


def _sheet(wb: Workbook, title: str, header: list[str], rows: list[list]):
    ws = wb.create_sheet(title)
    ws.append(header)
    for c in ws[1]:
        c.font = Font(bold=True)
    for r in rows:
        ws.append([("; ".join(map(str, v)) if isinstance(v, (list, tuple)) else v) for v in r])
    for i, h in enumerate(header, 1):
        ws.column_dimensions[get_column_letter(i)].width = min(max(12, len(str(h)) + 2), 48)
    ws.freeze_panes = "A2"
    return ws


def to_xlsx(result: dict, meta: dict) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    a = result.get("assumptions", {})
    info = [("Study", meta.get("id")), ("Kind", result.get("kind")), ("Project", meta.get("project_id")),
            ("Configuration version", meta.get("config_version_no")), ("Configuration hash", meta.get("config_hash")),
            ("Model version", result.get("model_version")), ("Climate dataset", (result.get("dataset") or {}).get("product")),
            ("Data nature", (result.get("dataset") or {}).get("data_nature")), ("Dataset SHA-256", meta.get("dataset_sha256")),
            ("Period start (UTC)", result.get("period", {}).get("start_utc")), ("Period end (UTC)", result.get("period", {}).get("end_utc")),
            ("Hours", result.get("period", {}).get("hours")), ("Simulations run", result.get("simulations_run")),
            ("Baseline", a.get("baseline")), ("Energy boundary", a.get("energy_boundary")), ("Water boundary", a.get("water_boundary")),
            ("Carbon boundary", a.get("carbon_boundary")), ("Plant sizing", a.get("plant_sizing")), ("Currency", a.get("currency")),
            ("Controller", result.get("controller_note"))]
    for k, v in info:
        ws.append([k, v])
    ws.append([])
    ws.append(["Limitations"])
    ws.cell(ws.max_row, 1).font = Font(bold=True)
    for lim in LIMITATIONS:
        ws.append(["", lim])
    ws.column_dimensions["A"].width = 26
    ws.column_dimensions["B"].width = 120

    rows = long_rows(result)
    if rows:
        _sheet(wb, "Metrics (long)", list(rows[0].keys()), [list(r.values()) for r in rows])
    sc = scenarios(result)
    keys = [r["key"] for r in result["baseline"]["metrics"]]
    labels = {r["key"]: (r["label"], r["unit"]) for r in result["baseline"]["metrics"]}
    wide = []
    for k in keys:
        row = [labels[k][0], labels[k][1]]
        for ev in sc:
            m = next((r for r in ev.get("metrics", []) if r["key"] == k), {})
            row += [m.get("value"), m.get("pct"), m.get("status")]
        wide.append(row)
    hdr = ["Metric", "Unit"] + [f"{ev['id']} {h}" for ev in sc for h in ("value", "% vs baseline", "status")]
    _sheet(wb, "Comparison (wide)", hdr, wide)

    cap = []
    for ev in sc:
        for ln in (ev.get("capex_incremental") or {}).get("lines", []):
            d = ln["delta_cost"] or {}
            cap.append([ev["id"], ln["label"], ln["unit"], ln["quantity_baseline"], ln["quantity_scenario"], ln["delta_quantity"],
                        d.get("low"), d.get("base"), d.get("high"), ln["scope"], ", ".join(ln["owners"]) if ln["owners"] else ""])
    _sheet(wb, "CAPEX changes", ["Scenario", "Line item", "Unit", "Qty baseline", "Qty scenario", "Δ qty", "Δ cost low", "Δ cost base",
                                 "Δ cost high", "Scope", "Used by"], cap)
    lines = []
    for ln in (result["baseline"].get("capex") or {}).get("lines", []):
        rate = ln.get("rate") or {}
        lines.append([ln["id"], ln["label"], ln["category"], ln["scope"], ln["quantity"], ln["unit"], ln["quantity_status"], ln["quantity_basis"],
                      rate.get("low"), rate.get("base"), rate.get("high"), rate.get("source"), rate.get("basis_year"), ln["status"],
                      ln["includes"], ln["excludes"]])
    _sheet(wb, "Baseline CAPEX lines", ["Id", "Line item", "Category", "Scope", "Quantity", "Unit", "Quantity status", "Quantity basis",
                                        "Rate low", "Rate base", "Rate high", "Rate source", "Basis year", "Status", "Includes", "Excludes"], lines)
    eco = []
    for ev in sc[1:]:
        e = ev.get("economics") or {}
        eco.append([ev["id"], e.get("annual_opex_baseline"), e.get("annual_opex_scenario"), e.get("annual_energy_water_saving"), e.get("npv"),
                    e.get("simple_payback_years"), e.get("discounted_payback_years"), e.get("note") or e.get("payback_note")])
    _sheet(wb, "OPEX & payback", ["Scenario", "Annual OPEX baseline", "Annual OPEX scenario", "Energy+water saving/yr", "NPV",
                                  "Simple payback (yr)", "Discounted payback (yr)", "Note"], eco)
    sp = []
    for ev in sc:
        for cat, c in ((ev.get("space") or {}).get("categories") or {}).items():
            sp.append([ev["id"], cat, c["base"], c["low"], c["high"], c["status"],
                       "; ".join(f"{i['label']}: {i['base']:.0f} m² ({i['status']}, {i['basis']})" for i in c.get("items", []))])
    _sheet(wb, "Space", ["Scenario", "Category", "m² base", "m² low", "m² high", "Status", "Items"], sp)
    if result.get("attribution"):
        at = result["attribution"]
        ids = list(at["rows"][0]["individual"].keys()) if at["rows"] else []
        _sheet(wb, "Attribution", ["Metric", "Unit"] + [f"Δ {i} alone" for i in ids] + ["Σ individual", "Δ combined", "Non-additive"],
               [[r["label"], r["unit"]] + [r["individual"].get(i) for i in ids] + [r["sum_individual"], r["combined"], r["non_additive"]]
                for r in at["rows"]])
    if result.get("pairs"):
        rows_i = []
        for p in result["pairs"]:
            for key, it in p.get("interactions", {}).items():
                rows_i.append([p["a"], p["b"], key, it.get("sum_individual"), it.get("combined"), it.get("interaction"), it.get("class")])
            if not p["simulated"]:
                rows_i.append([p["a"], p["b"], "—", None, None, None, "infeasible: " + "; ".join(p["infeasible_reasons"])])
        _sheet(wb, "Pairwise interactions", ["A", "B", "Metric", "Σ individual Δ", "Combined Δ", "Interaction", "Class"], rows_i)
    _sheet(wb, "Strategies", ["Id", "Name", "Category", "Inspiration", "Mechanism", "Physics model", "Requires", "Exclusive group",
                              "Model status", "Validation"],
           [[m.id, m.name, m.category, m.inspiration, m.mechanism, m.physics_model, ", ".join(m.requires), m.exclusive_group,
             m.model_status, m.validation] for m in catalogue.MEASURES])
    rates = a.get("unit_rates", {})
    _sheet(wb, "Assumptions", ["Item", "Value", "Note"],
           [[f"unit rate: {k}", f"{v.get('low')} / {v.get('base')} / {v.get('high')}", f"{v.get('source')}, basis {v.get('basis_year')}; {v.get('note')}"]
            for k, v in rates.items()]
           + [[f"layout: {k}", v, ""] for k, v in (a.get("layout") or {}).items()]
           + [["contingency fraction", a.get("contingency_fraction"), ""], ["excluded costs", "; ".join(EXCLUDED_GLOBAL), ""],
              ["space", a.get("space_note"), ""]])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
