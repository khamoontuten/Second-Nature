"""Excel workbook export (14 sheets). Key ratios are written as live Excel formulas next to the stored values,
so the engineer can audit them; every number comes from the stored simulation."""

from __future__ import annotations

import io
import json

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from app.reporting.data import LIMITATIONS, ReportData

HEAD = Font(bold=True, color="FFFFFF")
FILL = PatternFill("solid", fgColor="1F4E79")
NOTE = Font(italic=True, color="555555")


def _sheet(wb, title, header, rows, widths=None, note=None):
    ws = wb.create_sheet(title[:31])
    r0 = 1
    if note:
        ws.cell(1, 1, note).font = NOTE
        r0 = 3
    for j, h in enumerate(header, 1):
        c = ws.cell(r0, j, h)
        c.font, c.fill = HEAD, FILL
        c.alignment = Alignment(wrap_text=True, vertical="top")
    for i, row in enumerate(rows, r0 + 1):
        for j, v in enumerate(row, 1):
            ws.cell(i, j, _cell(v))
    for j in range(1, len(header) + 1):
        ws.column_dimensions[get_column_letter(j)].width = (widths or {}).get(j, 18)
    ws.freeze_panes = ws.cell(r0 + 1, 1)
    return ws, r0


def _cell(v):
    if isinstance(v, dict | list):
        return json.dumps(v, default=str)[:32000]
    if isinstance(v, float) and v != v:
        return None
    return v


def build_workbook(d: ReportData) -> bytes:
    wb = Workbook()
    wb.remove(wb.active)
    run, comp = d.run, d.run.comparison or {}
    clim = run.climate_summary or {}
    base = d.results[d.baseline_id].summary
    # 1 Project Inputs
    rows = [[a["path"], a["value"], a["source"], a.get("description", "")] for a in d.inputs.assumption_register]
    _sheet(wb, "Project Inputs", ["Parameter", "Value", "Source", "Description"], rows, {1: 46, 2: 28, 3: 18, 4: 60},
           note=f"Project {d.project.name} · config hash {run.config_hash[:12]} · model {run.model_version} · "
                "'default-assumption' = value not supplied by the engineer")
    # 2 Climate Summary
    rows = []
    for var, st in (clim.get("monthly") or {}).items():
        for stat, vals in st.items():
            rows.append([var, stat] + list(vals))
    ws, r0 = _sheet(wb, "Climate Summary", ["Variable", "Statistic"] + [f"M{m}" for m in range(1, 13)], rows,
                    note=f"Period {clim.get('period', {}).get('start_utc')} – {clim.get('period', {}).get('end_utc')}; "
                         f"nature: {clim.get('period', {}).get('data_nature')}")
    dc = clim.get("design_conditions", {})
    r = ws.max_row + 2
    ws.cell(r, 1, "Design conditions (dataset exceedance values)").font = Font(bold=True)
    for i, (v, vals) in enumerate(dc.items(), r + 1):
        ws.cell(i, 1, v)
        for j, (k, x) in enumerate(vals.items(), 2):
            ws.cell(i, j, f"{k}: {x:.2f}")
    # 3 Data Provenance
    prov = d.dataset.provenance
    _sheet(wb, "Data Provenance", ["Field", "Value"], [[k, v] for k, v in prov.items()], {1: 28, 2: 120},
           note=f"Dataset {d.dataset.id} · SHA-256 {d.dataset.sha256}")
    # 4 Free Cooling
    rows = []
    for sid in d.scenario_ids:
        fc = d.results[sid].summary["free_cooling"]
        el = fc["eligibility"]
        rows.append([sid, d.results[sid].label, fc["full_hours"], fc["partial_hours"], fc["mechanical_hours"],
                     fc["residual_mechanical_cooling_mwh"], el["dry_bulb_eligible_h"], el["enthalpy_eligible_h"],
                     el["full_envelope_eligible_h"], el["rejected_humidity_high_h"], el["rejected_humidity_low_h"],
                     el["rejected_temperature_h"], el["rejected_air_quality_h"]])
    _sheet(wb, "Free Cooling", ["Scenario", "Label", "Full free-cooling h", "Partial h", "Mechanical h", "Residual mech. cooling MWh",
                                "Dry-bulb eligible h", "Enthalpy eligible h", "Full-envelope eligible h", "Rejected: humidity high h",
                                "Rejected: humidity low h", "Rejected: temperature h", "Rejected: air quality h"], rows,
           note="Environmental eligibility ≠ useful cooling: full/partial hours come from the hourly controller with capacity and airflow checks.")
    # 5 Airflow Model
    rows = []
    for sid in d.scenario_ids:
        nv = d.results[sid].diagnostics.get("natural_ventilation")
        if nv:
            st = d.results[sid].config["strategies"]["stack"]
            rows.append([sid, st["chimney_height_m"], st["chimney_area_m2"], st["n_chimneys"], st["discharge_coefficient"],
                         nv["sufficient_hours"], nv["reversed_hours"], nv["mean_natural_flow_kg_s"], nv["mean_required_kg_s"],
                         nv["median_fan_dp_pa"], nv["network_converged_fraction"]])
    _sheet(wb, "Airflow Model", ["Scenario", "Chimney height m", "Chimney area m²", "Chimneys", "C_d", "Natural-sufficient h",
                                 "Reversed-flow h", "Mean natural flow kg/s", "Mean required kg/s", "Median assist Δp Pa",
                                 "Network converged fraction"], rows,
           note="ṁ = C_d A √(2ρΔp); Δp_stack = g H (ρ_out − ρ_exhaust); wind p = ½ρU²C_p; zone pressures solved by Newton–Raphson.")
    # 6 Earth-Air Heat Exchanger
    rows = []
    for sid in d.scenario_ids:
        e = d.results[sid].summary.get("eahe")
        if e:
            g = d.results[sid].config["strategies"]["eahe"]
            for y in e["yearly"]:
                rows.append([sid, g["n_pipes"], g["length_m"], g["diameter_m"], g["depth_m"], g["spacing_m"], y["year"],
                             y["sensible_cooling_kwh"] / 1000, y["sensible_heating_kwh"] / 1000, y["condensate_m3"], y["mean_drift_k"],
                             y["max_drift_k"], e["fan_energy_mwh"], e["useful_hours"]])
    _sheet(wb, "Earth-Air Heat Exchanger", ["Scenario", "Pipes", "Length m", "Diameter m", "Depth m", "Spacing m", "Year",
                                            "Sensible cooling MWh", "Sensible heating MWh", "Condensate m³", "Mean soil drift K",
                                            "Max soil drift K", "Fan MWh (final year)", "Useful h (final year)"], rows,
           note="T_out = T_w + (T_in − T_w) e^{−NTU}; radial implicit FD soil with shape-factor far field; years repeat the climate record.")
    # 7 Envelope Model
    rows = []
    for sid in d.scenario_ids:
        env = d.results[sid].diagnostics.get("envelope", {})
        for name, c in (env.get("constructions") or {}).items():
            rows.append([sid, name, c.get("name"), c.get("u_value"), c.get("areal_heat_capacity_j_m2k"), c.get("decrement_factor"),
                         c.get("time_lag_h"), c.get("solar_absorptance"), c.get("emissivity")])
    ws, _ = _sheet(wb, "Envelope Model", ["Scenario", "Element", "Construction", "U W/m²K", "Areal heat capacity J/m²K",
                                          "Decrement factor", "Time lag h", "Solar absorptance", "Emissivity"], rows)
    # 8 Cooling Equipment
    rows = []
    for sid in d.scenario_ids:
        sz = d.results[sid].diagnostics.get("sizing", {})
        comps = d.results[sid].summary["energy_mwh"]["components"]
        rows.append([sid, sz.get("architecture"), sz.get("q_design_kw"), sz.get("n_duty"), sz.get("n_installed"), sz.get("unit_capacity_kw"),
                     (sz.get("chiller") or {}).get("rated_cop"), sz.get("eta_II"), (sz.get("chiller") or {}).get("model"),
                     comps.get("p_chiller", 0), comps.get("p_crah_fans", 0), comps.get("p_chw_pumps", 0), comps.get("p_tower_fans", 0),
                     comps.get("p_econ_fans", 0), comps.get("p_fan_assist", 0), comps.get("p_eahe_fans", 0), comps.get("p_iec_fans", 0)])
    _sheet(wb, "Cooling Equipment", ["Scenario", "Architecture", "Design load kW", "Duty units", "Installed units", "Unit capacity kW",
                                     "Rated COP", "η_II (Carnot fraction)", "Chiller model", "Chiller MWh", "CRAH fans MWh",
                                     "CHW pumps MWh", "Tower fans MWh", "Economiser fans MWh", "Stack assist MWh", "EAHE fans MWh",
                                     "IEC fans MWh"], rows)
    # 9 Energy Balance (with formulas)
    ws = wb.create_sheet("Energy Balance")
    hdr = ["Scenario", "IT MWh", "Facility MWh", "Cooling MWh", "PUE (stored)", "PUE (formula)", "Mean hourly PUE",
           "Heat to plant condenser MWh", "Economiser exhaust MWh", "Ground MWh", "IEC secondary MWh", "Radiative MWh",
           "Closure max relative residual", "Closure passed"]
    for j, h in enumerate(hdr, 1):
        c = ws.cell(1, j, h)
        c.font, c.fill = HEAD, FILL
    for i, sid in enumerate(d.scenario_ids, 2):
        sm = d.results[sid].summary
        hp = sm["heat_mwh"]["paths"]
        vals = [sid, sm["energy_mwh"]["it"], sm["energy_mwh"]["facility"], sm["energy_mwh"]["cooling"], sm["pue"]["annual"], None,
                sm["pue"]["mean_of_hourly"], hp["plant_condenser"], hp["economiser_exhaust"], hp["ground_eahe"], hp["iec_secondary_air"],
                hp["radiative_sky"], sm["heat_mwh"]["closure"]["max_relative_residual"], sm["heat_mwh"]["closure"]["passed"]]
        for j, v in enumerate(vals, 1):
            ws.cell(i, j, v)
        ws.cell(i, 6, f"=C{i}/B{i}")
    for j in range(1, len(hdr) + 1):
        ws.column_dimensions[get_column_letter(j)].width = 18
    # 10 Water Balance
    ws = wb.create_sheet("Water Balance")
    hdr = ["Scenario", "IT MWh", "Withdrawal m³", "Consumption m³", "WUE withdrawal L/kWh (stored)", "WUE (formula)", "Potable m³",
           "Reclaimed m³", "Peak hourly m³", "Peak daily m³", "Tower evaporation m³", "Tower blowdown m³", "Evaporative media m³",
           "Humidification m³"]
    for j, h in enumerate(hdr, 1):
        c = ws.cell(1, j, h)
        c.font, c.fill = HEAD, FILL
    for i, sid in enumerate(d.scenario_ids, 2):
        w = d.results[sid].summary["water"]
        cm = w["components_m3"]
        vals = [sid, d.results[sid].summary["energy_mwh"]["it"], w["withdrawal_m3"], w["consumption_m3"], w["wue_withdrawal_l_per_kwh"],
                None, w["potable_m3"], w["reclaimed_m3"], w["peak_hourly_withdrawal_m3"], w["peak_daily_withdrawal_m3"],
                cm["wl_tower_evap"], cm["wl_tower_blowdown"], cm["wl_evap_media"], cm["wl_humidification"]]
        for j, v in enumerate(vals, 1):
            ws.cell(i, j, v)
        ws.cell(i, 6, f"=C{i}*1000/(B{i}*1000)")
    ws.cell(len(d.scenario_ids) + 3, 1, f"Boundary: {base['water']['boundary']}").font = NOTE
    # 11 Carbon Accounting
    ws = wb.create_sheet("Carbon Accounting")
    hdr = ["Scenario", "Facility MWh", "Operational t CO₂e (stored)", "Cooling-related t CO₂e", "Basis", "Factor type", "Source",
           "Marginal-factor scenario t"]
    for j, h in enumerate(hdr, 1):
        c = ws.cell(1, j, h)
        c.font, c.fill = HEAD, FILL
    for i, sid in enumerate(d.scenario_ids, 2):
        cb = d.results[sid].summary["carbon"]
        for j, v in enumerate([sid, d.results[sid].summary["energy_mwh"]["facility"], cb["operational_t"], cb["cooling_related_t"],
                               cb["intensity_basis"], cb["factor_type"], cb["source"], cb["marginal_scenario_t"]], 1):
            ws.cell(i, j, v)
    ws.cell(len(d.scenario_ids) + 3, 1, "Operational emissions only; embodied carbon is reported separately (lifecycle).").font = NOTE
    # 12 Scenario Comparison
    metrics = list(next(iter(comp.values()))["metrics"].keys()) if comp else []
    rows = []
    for sid in sorted(comp):
        row = comp[sid]
        rows.append([sid, row["label"], "+".join(row["strategies"]), row["objective"]] + [row["metrics"].get(m) for m in metrics]
                    + [row["delta_pct"].get(m) for m in metrics])
    _sheet(wb, "Scenario Comparison", ["Scenario", "Label", "Strategies", "Objective"] + metrics + [f"Δ% {m}" for m in metrics], rows,
           note="All scenarios share the same climate, IT load, envelope class and redundancy; Δ% relative to S0.")
    # 13 Sensitivity Analysis
    rows = []
    for sres in d.sensitivity:
        res = sres.result or {}
        if sres.kind == "tornado":
            for r in res.get("rows", []):
                rows.append(["tornado", r["label"], r["category"], r["low_value"], r["high_value"], r["low_saving"], r["high_saving"], None])
        elif sres.kind == "monte_carlo":
            for m, v in res.get("metrics", {}).items():
                sv = v.get("saving", {})
                rows.append(["monte_carlo", m, "all", sv.get("p5"), sv.get("p95"), sv.get("p50"), sv.get("mean"), v.get("prob_design_better")])
    _sheet(wb, "Sensitivity Analysis", ["Method", "Input / metric", "Category", "Low / P5", "High / P95", "Saving low / P50",
                                        "Saving high / mean", "P(design better)"], rows,
           note="Run tornado / Monte Carlo from the Optimisation page to populate this sheet." if not rows else None)
    # 14 Assumptions and Limitations
    rows = [["limitation", x] for x in LIMITATIONS]
    for sid in d.scenario_ids:
        rows += [[f"{sid} flag ({f['level']})", f["message"]] for f in d.results[sid].flags]
    defaults = [a for a in d.inputs.assumption_register if a["source"] == "default-assumption"]
    rows.append(["data gaps", f"{len(defaults)} parameters use default assumptions (see Project Inputs)"])
    _sheet(wb, "Assumptions and Limitations", ["Type", "Statement"], rows, {1: 28, 2: 140})
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
