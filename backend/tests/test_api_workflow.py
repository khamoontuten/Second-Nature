"""API integration and end-to-end workflow (§22.4/22.5): create → site → climate → configure → simulate →
compare → optimise → inspect → export → reload → reproduce. External services are mocked."""

import io
import json
import time
import zipfile

import httpx
import openpyxl
import pytest
import respx
from fastapi.testclient import TestClient

from app.main import app


@pytest.fixture(scope="module")
def client():
    with TestClient(app) as c:
        yield c


def wait(client, job_id, timeout=900):
    t0 = time.time()
    while time.time() - t0 < timeout:
        j = client.get(f"/api/jobs/{job_id}").json()
        if j["status"] in ("succeeded", "failed", "cancelled"):
            return j
        time.sleep(0.3)
    raise AssertionError("job timeout")


CONFIG = {"name": "API test — Phoenix", "site": {"name": "Phoenix test", "latitude": 33.45, "longitude": -111.98, "elevation_m": 337,
                                                 "timezone": "America/Phoenix"},
          "grid": {"emission_factor_kg_per_kwh": 0.4, "source": "test"},
          "economics": {"electricity_price_per_kwh": 0.1, "water_price_per_m3": 2.0}}


@pytest.fixture(scope="module")
def project(client):
    r = client.post("/api/projects", json={"config": CONFIG, "reference_site_id": "phoenix_az"})
    assert r.status_code == 201, r.text
    return r.json()


def test_health_and_reference(client):
    h = client.get("/api/health").json()
    assert h["status"] == "ok" and h["database"]
    assert len(client.get("/api/strategies").json()["strategies"]) == 9
    assert any(e.get("class") == "A1" for e in client.get("/api/environmental-classes").json()["envelopes"])
    assert len(client.get("/api/climate/reference-sites").json()["sites"]) == 8
    assert {p["id"] for p in client.get("/api/climate/providers").json()} >= {"open_meteo", "nasa_power", "cds_era5"}


def test_project_create_validate_patch_versions(client, project):
    assert project["default_dataset_id"]
    assert any(a["source"] == "default-assumption" for a in project["assumptions"])
    assert any(a["path"] == "site.latitude" and a["source"] == "user" for a in project["assumptions"])
    r = client.patch(f"/api/projects/{project['id']}", json={"patch": {"it": {"installed_capacity_mw": 4.0}}, "note": "smaller"})
    assert r.status_code == 200 and r.json()["current_version_no"] == 2
    v1 = client.get(f"/api/projects/{project['id']}/versions/1").json()
    assert v1["config"]["it"]["installed_capacity_mw"] == 5.0  # earlier version preserved
    bad = client.post("/api/projects", json={"config": {"site": {"latitude": 200, "longitude": 0}}})
    assert bad.status_code == 422 and bad.json()["error"] == "validation_failed"
    assert client.get("/api/projects/doesnotexist").status_code == 404


def test_climate_summary_timeseries_and_upload(client, project):
    ds = project["default_dataset_id"]
    s = client.get(f"/api/climate/{ds}/summary", params={"timezone": "America/Phoenix", "elevation_m": 337}).json()
    assert s["summary"]["overall"]["t_db_c"]["max"] > 40 and s["provenance"]["data_nature"] == "typical_year"
    assert s["quality"]["timezone_check"]["plausible"]
    ts = client.get(f"/api/climate/{ds}/timeseries", params={"variables": "t_db_c,t_wb_c", "step": 24}).json()
    assert len(ts["index"]) == 365 and "t_wb_c" in ts["columns"]
    csv = client.get(f"/api/climate/{ds}/timeseries", params={"format": "csv", "step": 24})
    assert csv.headers["content-type"].startswith("text/csv")
    body = "time,temperature,rh\n" + "\n".join(f"2023-01-{d:02d} {h:02d}:00,{15 + h % 10},{50}" for d in range(1, 31) for h in range(24))
    up = client.post("/api/climate/upload", files={"file": ("obs.csv", body.encode(), "text/csv")},
                     data={"latitude": 33.45, "longitude": -111.98, "project_id": project["id"], "description": "station test"})
    assert up.status_code == 201, up.text
    assert up.json()["quality"]["warnings"]  # < 1 year coverage flagged
    bad = client.post("/api/climate/upload", files={"file": ("x.exe", b"MZ...", "application/octet-stream")},
                      data={"latitude": 0, "longitude": 0})
    assert bad.status_code == 422
    fake_nc = client.post("/api/climate/upload", files={"file": ("x.nc", b"not netcdf", "application/x-netcdf")},
                          data={"latitude": 0, "longitude": 0})
    assert fake_nc.status_code == 422 and "NetCDF" in fake_nc.json()["message"]
    # reattach the reference dataset as default for later tests
    client.patch(f"/api/projects/{project['id']}", json={"patch": {}, "default_dataset_id": ds})


@respx.mock
def test_geocode_fallback_is_explicit(client):
    respx.get("https://geocoding-api.open-meteo.com/v1/search").mock(return_value=httpx.Response(503))
    respx.get(url__startswith="https://nominatim.openstreetmap.org").mock(return_value=httpx.Response(200, json=[]))
    r = client.post("/api/sites/geocode", json={"query": "Nowhere place xyz"}).json()
    assert r["results"] == [] and "manually" in r["fallback"]


@respx.mock
def test_climate_fetch_job_failure_is_reported(client):
    respx.get("https://archive-api.open-meteo.com/v1/archive").mock(return_value=httpx.Response(500))
    respx.get("https://power.larc.nasa.gov/api/temporal/hourly/point").mock(return_value=httpx.Response(503))
    r = client.post("/api/climate/fetch", json={"latitude": 12.97, "longitude": 77.59, "start": "2023-01-01", "end": "2023-01-02",
                                                "providers": ["open_meteo", "nasa_power"]})
    j = wait(client, r.json()["job_id"], timeout=120)
    assert j["status"] == "failed" and "all providers failed" in j["message"]


@pytest.fixture(scope="module")
def simulation(client, project):
    r = client.post("/api/simulations", json={"project_id": project["id"], "scenarios": ["S0", "S1", "S3"]})
    assert r.status_code == 202, r.text
    j = wait(client, r.json()["job_id"])
    assert j["status"] == "succeeded", j
    return r.json()["simulation_id"]


def test_simulation_results_and_hourly(client, simulation):
    res = client.get(f"/api/simulations/{simulation}/results").json()
    assert res["status"] == "succeeded" and set(res["comparison"]) == {"S0", "S1", "S3"}
    s0 = res["comparison"]["S0"]["metrics"]
    assert s0["unmet_hours"] == 0 and 1.0 < s0["pue"] < 2.0
    for sr in res["scenario_results"]:
        assert sr["summary"]["heat_mwh"]["closure"]["passed"]
    h = client.get(f"/api/simulations/{simulation}/scenarios/S3/hourly", params={"columns": "mode,p_facility_w,t_inlet_c", "step": 6}).json()
    assert len(h["index"]) == 1460 and set(h["columns"]) == {"mode", "p_facility_w", "t_inlet_c"}
    csv = client.get(f"/api/simulations/{simulation}/scenarios/S0/hourly", params={"format": "csv"})
    assert csv.status_code == 200 and csv.text.count("\n") > 8000
    assert client.get(f"/api/simulations/{simulation}/scenarios/S9/hourly").status_code == 404


def test_reproduce_from_saved_inputs(client, simulation):
    r = client.post(f"/api/simulations/{simulation}/reproduce", params={"scenario_id": "S1"}).json()
    assert r["reproduced"] and r["hourly_identical"]


def test_reproduce_run_stored_before_schema_fields_were_added(client, simulation):
    """A run saved before the configuration gained fields still reproduces (identical hourly results), with a note;
    a changed stored input does not."""
    from app.database.models import ScenarioResult
    from app.database.session import session_scope

    with session_scope() as s:
        row = s.query(ScenarioResult).filter_by(run_id=simulation, scenario_id="S0").one()
        original = (dict(row.config), row.inputs_hash)
        cfg = dict(row.config)
        cfg.pop("layout")
        cfg["economics"] = {k: v for k, v in cfg["economics"].items() if k not in ("unit_rates", "contingency_fraction", "capex_basis_note")}
        row.config, row.inputs_hash = cfg, "hash-from-older-schema"
    try:
        r = client.post(f"/api/simulations/{simulation}/reproduce", params={"scenario_id": "S0"}).json()
        assert r["hourly_identical"] and r["reproduced"] and "schema" in r["note"]
        with session_scope() as s:
            row = s.query(ScenarioResult).filter_by(run_id=simulation, scenario_id="S0").one()
            cfg = dict(row.config)
            cfg["it"] = {**cfg["it"], "average_utilisation": 0.5}
            row.config = cfg
        r = client.post(f"/api/simulations/{simulation}/reproduce", params={"scenario_id": "S0"}).json()
        assert not r["reproduced"] and not r["hourly_identical"]
    finally:
        with session_scope() as s:
            row = s.query(ScenarioResult).filter_by(run_id=simulation, scenario_id="S0").one()
            row.config, row.inputs_hash = original


def test_optimisation_and_uncertainty_jobs(client, project, simulation):
    o = client.post("/api/optimizations", json={"project_id": project["id"], "mode": "pareto", "strategies": ["stack"], "budget": 4,
                                               "variables": [{"path": "strategies.stack.chimney_height_m", "values": [10, 30]},
                                                             {"path": "strategies.stack.chimney_area_m2", "values": [15, 45]}]}).json()
    assert wait(client, o["job_id"])["status"] == "succeeded"
    res = client.get(f"/api/optimizations/{o['optimization_id']}").json()["result"]
    assert res["best"] and len(res["evaluations"]) == 4 and res["pareto_indices"]
    t = client.post("/api/uncertainty", json={"simulation_id": simulation, "kind": "tornado", "scenario_id": "S3"}).json()
    assert wait(client, t["job_id"])["status"] == "succeeded"
    rows = client.get(f"/api/uncertainty/{t['sensitivity_id']}").json()["result"]["rows"]
    assert rows and all("swing_saving" in r for r in rows)


def test_reports_exports_match_stored_values(client, simulation):
    res = client.get(f"/api/simulations/{simulation}/results").json()
    stored = {r["scenario_id"]: r["summary"] for r in res["scenario_results"]}
    files = {}
    for fmt in ("pdf", "xlsx", "json", "zip"):
        r = client.post("/api/reports", json={"simulation_id": simulation, "format": fmt}).json()
        j = wait(client, r["job_id"])
        assert j["status"] == "succeeded", j
        meta = client.get(f"/api/reports/{r['report_id']}").json()
        dl = client.get(meta["download_url"])
        assert dl.status_code == 200
        files[fmt] = dl.content
    assert files["pdf"][:4] == b"%PDF"
    wb = openpyxl.load_workbook(io.BytesIO(files["xlsx"]))
    assert len(wb.sheetnames) == 14
    eb = wb["Energy Balance"]
    rows = {eb.cell(i, 1).value: eb.cell(i, 3).value for i in range(2, eb.max_row + 1)}
    for sid, sm in stored.items():
        assert rows[sid] == pytest.approx(sm["energy_mwh"]["facility"], rel=1e-12)
    bundle = json.loads(files["json"])
    for sid, sm in stored.items():
        assert bundle["scenarios"][sid]["summary"]["energy_mwh"]["facility"] == pytest.approx(sm["energy_mwh"]["facility"])
    z = zipfile.ZipFile(io.BytesIO(files["zip"]))
    assert "scenario_comparison.csv" in z.namelist() and "hourly_S0.csv" in z.namelist()
    import pandas as pd

    h0 = pd.read_csv(z.open("hourly_S0_energy_kwh.csv"))
    assert h0["p_facility_w"].sum() / 1000 == pytest.approx(stored["S0"]["energy_mwh"]["facility"], rel=1e-9)


def test_project_reload_lists_runs(client, project, simulation):
    p = client.get(f"/api/projects/{project['id']}").json()
    assert any(s["id"] == simulation for s in p["simulations"])
    assert p["versions"][-1]["version_no"] >= 2


def test_calculators(client):
    ps = client.post("/api/calculators/psychrometrics", json={"t_db_c": 30, "rh_pct": 50}).json()
    assert ps["t_wb_c"] == pytest.approx(22.0, abs=0.1)
    st = client.post("/api/calculators/stack", json={"heights_m": [5, 20, 40], "required_flow_m3s": 100}).json()
    flows = [r["volume_flow_m3s"] for r in st["rows"]]
    assert flows == sorted(flows) and st["warning"]
    ea = client.post("/api/calculators/earth-air", json={"hours": 48, "sweep": "length"}).json()
    assert ea["last_hour"]["t_out_c"] < 32 and len(ea["sweep"]["rows"]) == 6
    en = client.post("/api/calculators/envelope", json={"layers": [{"material": "concrete_dense", "thickness_m": 0.2},
                                                                   {"material": "mineral_wool", "thickness_m": 0.1}]}).json()
    assert 0 < en["decrement_factor"] < 1 and en["time_lag_h"] > 3
    af = client.post("/api/calculators/airflow-requirement", json={"it_load_kw": 1000, "delta_t_k": 12}).json()
    assert af["mass_flow_kg_s"] == pytest.approx(1e6 / ((1006 + 1860 * 0.009) * 12))
    bad = client.post("/api/calculators/psychrometrics", json={"t_db_c": 20, "t_dp_c": 25})
    assert bad.status_code == 422


def test_job_cancel_and_unknown(client):
    assert client.get("/api/jobs/nope").status_code == 404
    assert client.post("/api/jobs/nope/cancel").json()["cancel_requested"] is False
